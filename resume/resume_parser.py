
"""Extract readable text from PDF, DOCX, and TXT resumes."""

from io import BytesIO
from pathlib import Path

from werkzeug.utils import secure_filename


ALLOWED_RESUME_EXTENSIONS = {"pdf", "docx", "txt"}
MAX_RESUME_CHARS = 30000
MAX_RESUME_BYTES = 8 * 1024 * 1024


class ResumeParseError(ValueError):
    """Raised when a resume cannot be parsed safely."""


def _get_extension(filename):
    return Path(secure_filename(filename or "")).suffix.lower().lstrip(".")


def extract_resume_text(source, filename=None):
    """
    Extract text from a path, bytes object, file-like object,
    or Flask FileStorage.
    """
    raw_bytes = None

    if hasattr(source, "read"):
        filename = filename or getattr(source, "filename", None)

        try:
            current_pos = source.tell()
        except Exception:
            current_pos = None

        raw_bytes = source.read()

        if current_pos is not None:
            try:
                source.seek(current_pos)
            except Exception:
                pass

    elif isinstance(source, (bytes, bytearray)):
        raw_bytes = bytes(source)

    else:
        path = Path(source)
        filename = filename or path.name

        if path.suffix.lower() == ".txt":
            try:
                text = path.read_text(encoding="utf-8", errors="replace")
            except OSError as exc:
                raise ResumeParseError(
                    f"Could not read TXT resume: {exc}"
                ) from exc

            text = text.replace("\x00", "")
            if len(text.strip()) < 30:
                raise ResumeParseError(
                    "Very little text was extracted from the resume."
                )

            return text[:MAX_RESUME_CHARS]

        try:
            raw_bytes = path.read_bytes()
        except OSError as exc:
            raise ResumeParseError(
                f"Could not read resume file: {exc}"
            ) from exc

    if not filename:
        raise ResumeParseError("Resume filename is missing.")

    safe_name = secure_filename(filename)
    extension = _get_extension(safe_name)

    if extension not in ALLOWED_RESUME_EXTENSIONS:
        raise ResumeParseError(
            "Unsupported resume format. Upload PDF, DOCX, or TXT."
        )

    if not raw_bytes:
        raise ResumeParseError("The uploaded resume is empty.")

    if len(raw_bytes) > MAX_RESUME_BYTES:
        raise ResumeParseError("Resume exceeds the 8 MB size limit.")

    if extension == "txt":
        text = raw_bytes.decode("utf-8", errors="replace")

    elif extension == "pdf":
        try:
            try:
                import pymupdf
                open_pdf = pymupdf.open
            except ImportError:
                import fitz
                open_pdf = fitz.open

            with open_pdf(stream=raw_bytes, filetype="pdf") as document:
                text = "\n".join(
                    page.get_text("text") for page in document
                )

        except Exception as exc:
            raise ResumeParseError(
                f"Could not read PDF: {exc}"
            ) from exc

    else:  # DOCX
        try:
            from docx import Document

            document = Document(BytesIO(raw_bytes))
            parts = [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]

            for table in document.tables:
                for row in table.rows:
                    parts.append(
                        " | ".join(cell.text.strip() for cell in row.cells)
                    )

            text = "\n".join(parts)

        except Exception as exc:
            raise ResumeParseError(
                f"Could not read DOCX: {exc}"
            ) from exc

    text = text.replace("\x00", "")
    text = "\n".join(
        line.strip()
        for line in text.splitlines()
        if line.strip()
    )

    if len(text.strip()) < 30:
        raise ResumeParseError(
            "Very little text was extracted. If this is a scanned PDF, "
            "run OCR first or upload a text-based PDF."
        )

    return text[:MAX_RESUME_CHARS]


def parse_resume(source, filename=None):
    """Return parsed text and file metadata."""
    name = filename or getattr(source, "filename", None)

    if name is None and not isinstance(source, (bytes, bytearray)):
        name = Path(source).name

    if not name:
        raise ResumeParseError("Resume filename is missing.")

    text = extract_resume_text(source, name)
    safe_name = secure_filename(name)

    return {
        "filename": safe_name,
        "file_type": _get_extension(safe_name),
        "text": text,
        "character_count": len(text),
        "word_count": len(text.split()),
    }
