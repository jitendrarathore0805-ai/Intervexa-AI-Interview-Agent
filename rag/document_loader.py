
"""Extract text from PDF, DOCX, and TXT files."""

from pathlib import Path


def extract_pdf_text(file_path):
    """Extract text from a PDF file."""
    try:
        import pymupdf

        with pymupdf.open(file_path) as pdf:
            text = "\n".join(page.get_text("text") for page in pdf)

        return text.strip()

    except Exception as exc:
        raise ValueError(f"Could not extract PDF text: {exc}") from exc


def extract_docx_text(file_path):
    """Extract text from DOCX paragraphs and tables."""
    try:
        from docx import Document

        document = Document(file_path)
        parts = []

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text.strip())

        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]
                if cells:
                    parts.append(" | ".join(cells))

        return "\n".join(parts).strip()

    except Exception as exc:
        raise ValueError(f"Could not extract DOCX text: {exc}") from exc


def extract_text(file_path):
    """
    Extract text from a PDF, DOCX, or TXT file path.

    This function accepts a filesystem path, not Flask FileStorage.
    """
    path = Path(file_path)
    extension = path.suffix.lower()

    if not path.is_file():
        raise FileNotFoundError(f"File not found: {path}")

    if extension == ".pdf":
        text = extract_pdf_text(path)

    elif extension == ".docx":
        text = extract_docx_text(path)

    elif extension == ".txt":
        try:
            text = path.read_text(encoding="utf-8", errors="replace")
        except OSError as exc:
            raise ValueError(f"Could not read TXT file: {exc}") from exc

    else:
        raise ValueError(
            "Unsupported file format. Use PDF, DOCX or TXT."
        )

    if not text.strip():
        raise ValueError(
            "No readable text found. Scanned PDFs may require OCR."
        )

    return text.strip()
